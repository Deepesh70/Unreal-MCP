from pptx import Presentation
p=Presentation(r'C:\Users\aadit\Downloads\SIH2026-IDEA-Presentation-Format.pptx')
print('slides',len(p.slides))
for i,s in enumerate(p.slides,1):
 print(i, [sh.text for sh in s.shapes if hasattr(sh,'text_frame')])
